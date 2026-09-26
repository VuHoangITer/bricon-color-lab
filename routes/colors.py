import re as _re
import time
from io import BytesIO
from datetime import datetime
from pathlib import Path
from flask import (Blueprint, render_template, request, redirect, url_for,
                   flash, session, g, abort, send_file)
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import config
from models import color as color_model
from models import color_version
from models import pigment as pigment_model
from models import image as image_model
from models.activity_log import log
from routes.auth import login_required, permission_required
from services.permissions import ROLE_ADMIN

bp = Blueprint("colors", __name__)


def _clean_hex(value):
    v = (value or "").strip().upper()
    return v if _re.fullmatch(r"#[0-9A-F]{6}", v) else None


def _parse_lab(form, prefix=""):
    """Đọc 3 ô L*/a*/b* (đo từ máy đo màu) từ form — trả về tuple (L,a,b)
    hoặc None nếu cả 3 ô đều trống (chưa đo bằng máy). Bắt buộc nhập ĐỦ cả
    3 nếu có nhập, tránh lưu Lab nửa vời gây sai khi so màu."""
    raw_l = (form.get(f"{prefix}lab_l") or "").strip().replace(",", ".")
    raw_a = (form.get(f"{prefix}lab_a") or "").strip().replace(",", ".")
    raw_b = (form.get(f"{prefix}lab_b") or "").strip().replace(",", ".")
    if not raw_l and not raw_a and not raw_b:
        return None
    if not (raw_l and raw_a and raw_b):
        raise ValueError("Nhập đủ cả 3 số L*/a*/b* từ máy đo màu, hoặc để trống cả 3 nếu chưa đo.")
    try:
        l, a, b = float(raw_l), float(raw_a), float(raw_b)
    except ValueError:
        raise ValueError("Giá trị L*/a*/b* không hợp lệ — phải là số.")
    if not (0 <= l <= 100):
        raise ValueError("L* phải trong khoảng 0–100.")
    return (l, a, b)


def _parse_ratios(form, pigments):
    """Đọc input dạng % (0-100) từ form, trả về dict {pigment_id: phân số 0-1}."""
    ratios = {}
    for p in pigments:
        raw = (form.get(f"pigment_{p['id']}") or "").strip().replace(",", ".")
        try:
            val = float(raw) if raw else 0.0
        except ValueError:
            raise ValueError(f"Tỷ lệ của '{p['name']}' không hợp lệ: {raw}")
        if val < 0 or val > 100:
            raise ValueError(f"Tỷ lệ của '{p['name']}' phải trong khoảng 0–100%.")
        ratios[p["id"]] = val / 100.0
    return ratios


def _is_admin():
    """True nếu người thao tác là role admin — CHỈ admin được tự tay gõ/sửa
    mã màu. Các role khác luôn bị khóa: khi tạo màu thì dùng đúng mã được
    gợi ý tự động, khi sửa màu thì giữ nguyên mã cũ, bất kể form gửi gì lên."""
    user = g.get("current_user")
    return bool(user and user["role"] == ROLE_ADMIN)


def _can_self_approve():
    """True nếu người thao tác có quyền duyệt sẵn -> tự động ĐÃ DUYỆT,
    không phải qua bước CHỜ DUYỆT."""
    return "duyet_mau" in g.get("current_permissions", set())


def _suggested_ma_mau():
    """Khách hàng (role khach_hang, kể cả được cấp thêm quyền tao_mau) ->
    gợi ý dạng KH-{username}-NN, tách riêng khỏi dãy số nội bộ của công ty.
    Nội bộ (admin/quản lý/sản xuất) -> vẫn dùng dãy số như cũ (max+1)."""
    user = g.get("current_user")
    if user and user["role"] == "khach_hang":
        return color_model.next_ma_mau_khach_hang(user["username"])
    return color_model.next_ma_mau()


def _can_create_version(color):
    """Ai được tạo phiên bản mới cho màu này: có quyền sua_mau chung, hoặc
    là người tạo ra chính màu đó (kể cả khách hàng đặc biệt được cấp tao_mau)."""
    if "sua_mau" in g.get("current_permissions", set()):
        return True
    return color["created_by"] == session.get("user_id")


def _can_manage_images(color):
    return "upload_anh" in g.get("current_permissions", set()) or _can_create_version(color)


def _can_delete_color(color):
    """CHỈ admin được xóa màu — không còn cho Quản lý hay người tạo tự xóa
    nữa, bất kể màu đó đã có phiên bản dùng hay chưa."""
    return _is_admin()


def _can_delete_version(version):
    """duyet_mau (Admin/Quản lý) -> xóa được MỌI phiên bản (kể cả đã duyệt/
    đang dùng). Người tạo phiên bản chỉ xóa được bản nháp của mình (CHỜ
    DUYỆT/TỪ CHỐI) — không tự xóa được bản ĐÃ DUYỆT."""
    if "duyet_mau" in g.get("current_permissions", set()):
        return True
    uid = session.get("user_id")
    return version["created_by"] == uid and version["trang_thai_duyet"] != color_version.DUYET_DA


def _can_activate():
    """Chỉ Admin/Quản lý được chuyển đổi bản đang dùng — thay đổi này ảnh
    hưởng tới sản xuất thật nên cần người có quyền duyệt quyết định."""
    return "duyet_mau" in g.get("current_permissions", set())


def _can_edit_version(version):
    """Sửa TRỰC TIẾP 1 phiên bản (không tạo bản mới):
    - Admin/Quản lý (duyet_mau) -> sửa được BẤT KỲ phiên bản nào, kể cả
      đã duyệt/đang dùng — áp dụng ngay, không cần duyệt lại (họ vốn có
      quyền tự duyệt).
    - Người khác -> sửa được phiên bản CHỜ DUYỆT hoặc TỪ CHỐI do CHÍNH họ
      tạo (bản nháp chưa ai duyệt, hoặc bản bị từ chối — sửa xong sẽ tự
      gửi lại xin duyệt). Phiên bản đã duyệt/đang dùng thì không sửa tại
      chỗ được — phải tạo phiên bản mới."""
    if "duyet_mau" in g.get("current_permissions", set()):
        return True
    uid = session.get("user_id")
    return version["created_by"] == uid and version["trang_thai_duyet"] in (
        color_version.DUYET_CHO, color_version.DUYET_TU_CHOI)


@bp.route("/")
@login_required
def index():
    return render_template("colors/list.html",
                           items=color_model.list_all(),
                           pigments=pigment_model.all_ordered())


def _is_dark_hex(hex6):
    """True nếu màu tối -> dùng chữ trắng cho dễ đọc khi tô nền ô Excel
    đúng theo màu HEX thật."""
    r, g_, b = int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16)
    return (0.299 * r + 0.587 * g_ + 0.114 * b) < 140


@bp.route("/colors/export")
@login_required
def export_excel():
    """Xuất nhanh mã màu + mã HEX + tỷ lệ pigment cấu tạo (theo phiên bản
    đang dùng) ra file Excel — đúng danh sách màu đang hiện ở Thư viện màu.
    Ô mã HEX được tô luôn theo đúng màu thật cho dễ nhìn/đối chiếu."""
    items = color_model.list_all()
    pigments = pigment_model.all_ordered()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mã màu"

    headers = ["Mã màu", "Tên màu", "Mã HEX"] + [p["name"] for p in pigments]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, it in enumerate(items, start=1):
        c = it["color"]
        hex_color = (c["hex_color"] or "").strip()
        row_idx = i + 1
        ratio_map = {}
        if c["active_version_id"]:
            ratio_map = {r["pigment_id"]: r["ty_le"] for r in color_version.get_ratio_rows(c["active_version_id"])}

        ws.append([c["ma_mau"], c["ten_mau"] or "", hex_color]
                  + [ratio_map.get(p["id"]) or None for p in pigments])

        fill_hex = hex_color.lstrip("#").upper()
        if len(fill_hex) == 6:
            hex_cell = ws.cell(row=row_idx, column=3)
            hex_cell.fill = PatternFill("solid", fgColor=fill_hex)
            hex_cell.font = Font(color="FFFFFF" if _is_dark_hex(fill_hex) else "000000", bold=True)
            hex_cell.alignment = Alignment(horizontal="center")

        for p_idx in range(len(pigments)):
            cell = ws.cell(row=row_idx, column=4 + p_idx)
            cell.number_format = "0.00%"
            cell.alignment = Alignment(horizontal="center")

    widths = [16, 28, 12] + [14] * len(pigments)
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(items) + 1}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"ma_mau_bricon_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/colors/tim-mau-giong")
@login_required
def similar_search():
    hex_query = (request.args.get("hex") or "").strip()
    raw_l = (request.args.get("lab_l") or "").strip()
    raw_a = (request.args.get("lab_a") or "").strip()
    raw_b = (request.args.get("lab_b") or "").strip()
    lab_query = None
    if raw_l and raw_a and raw_b:
        try:
            lab_query = (float(raw_l.replace(",", ".")), float(raw_a.replace(",", ".")), float(raw_b.replace(",", ".")))
        except ValueError:
            flash("Giá trị L*/a*/b* không hợp lệ.", "error")

    results = []
    if lab_query:
        # Có Lab đo máy -> dùng trực tiếp, chính xác hơn suy từ HEX.
        results = color_model.find_similar(lab=lab_query, limit=24)
    elif hex_query:
        results = color_model.find_similar(hex_color=hex_query, limit=24)
    return render_template("colors/similar_search.html", hex_query=hex_query,
                           raw_l=raw_l, raw_a=raw_a, raw_b=raw_b,
                           lab_query=lab_query, results=results)


@bp.route("/colors/cua-toi")
@login_required
def mine():
    return render_template("colors/mine.html", items=color_version.list_mine(session["user_id"]))


@bp.route("/colors/cho-duyet")
@permission_required("duyet_mau")
def pending():
    return render_template("colors/pending.html", items=color_version.list_pending())


@bp.route("/colors/new", methods=["GET", "POST"])
@permission_required("tao_mau")
def new():
    pigments = pigment_model.all_ordered()
    if request.method == "POST":
        try:
            ratios = _parse_ratios(request.form, pigments)
            lab = _parse_lab(request.form)

            # Bắt buộc phải có ít nhất 1 file khi tạo màu mới (chọn nhiều
            # file 1 lần thì tất cả cùng gắn "loại ảnh"/ghi chú đã chọn).
            files = [f for f in request.files.getlist("image") if f and f.filename]
            if not files:
                raise ValueError("Cần tải lên ít nhất 1 file khi tạo màu mới.")
            for f in files:
                ext = f.filename.rsplit(".", 1)[-1].lower() if "." in f.filename else ""
                if not ext or ext in config.BLOCKED_UPLOAD_EXT:
                    raise ValueError(f"File '{f.filename}': đuôi .{ext or '?'} không được phép tải lên.")

            auto_approve = _can_self_approve()
            # Chỉ admin được gõ tay mã màu — các role khác luôn bị khóa,
            # LUÔN lấy lại mã gợi ý tự động ngay lúc lưu (không tin giá trị
            # form gửi lên, tránh mã cũ/bị sửa tay qua devtools).
            ma_mau = request.form.get("ma_mau", "").strip() if _is_admin() else _suggested_ma_mau()
            color_id = color_model.create(
                ma_mau,
                request.form.get("ten_mau", ""),
                request.form.get("ghi_chu", ""),
                session["user_id"],
                hex_color=_clean_hex(request.form.get("hex_color")),
                lab=lab,
            )
            color_version.create_version(
                color_id, ratios, session["user_id"],
                ghi_chu=request.form.get("version_ghi_chu", ""),
                auto_approve=auto_approve,
            )

            config.IMAGE_FOLDER.mkdir(parents=True, exist_ok=True)
            for i, f in enumerate(files):
                ext = f.filename.rsplit(".", 1)[-1].lower() if "." in f.filename else ""
                filename = f"color{color_id}_{int(time.time()*1000)}_{i}_{secure_filename(Path(f.filename).stem) or 'anh'}.{ext}"
                f.save(config.IMAGE_FOLDER / filename)
                image_model.create(color_id, f"uploads/images/{filename}",
                                   request.form.get("loai_anh", "Khác"),
                                   session["user_id"],
                                   request.form.get("anh_ghi_chu", "").strip())

            log(session["user_id"], "TẠO MÀU", "color", color_id,
                {"ma_mau": ma_mau, "phien_ban": 1,
                 "trang_thai_duyet": "DA_DUYET" if auto_approve else "CHO_DUYET"})
            if auto_approve:
                flash("Đã tạo màu mới (tự động duyệt).", "success")
            else:
                flash("Đã tạo màu mới — phiên bản 1 đang CHỜ DUYỆT, cần Quản lý/Admin duyệt trước khi cân.", "success")
            return redirect(url_for("colors.detail", color_id=color_id))
        except ValueError as e:
            flash(str(e), "error")
        except Exception as e:
            flash(f"Lỗi: {e}", "error")
    return render_template("colors/form.html", pigments=pigments,
                           image_types=config.IMAGE_TYPES,
                           suggested_ma_mau=_suggested_ma_mau())


@bp.route("/colors/<int:color_id>")
@login_required
def detail(color_id):
    c = color_model.get(color_id)
    if not c:
        flash("Không tìm thấy màu này.", "error")
        return redirect(url_for("colors.index"))
    active = color_version.get_active(color_id)
    active_detail = color_version.detail(active["id"]) if active else None
    versions = color_version.list_for_color(color_id)
    target_lab = color_model.get_lab(c)
    similar_colors = color_model.find_similar(lab=target_lab, exclude_id=color_id, limit=6) if target_lab else []
    return render_template("colors/detail.html", color=c,
                           active=active, active_detail=active_detail,
                           versions=versions,
                           images=image_model.list_for_color(color_id),
                           image_types=config.IMAGE_TYPES,
                           ready_for_weighing=color_model.is_ready_for_weighing(color_id),
                           can_create_version=_can_create_version(c),
                           can_delete_color=_can_delete_color(c),
                           can_activate=_can_activate(),
                           version_has_pending=color_version.has_pending(color_id),
                           similar_colors=similar_colors,
                           can_edit_version=_can_edit_version)


@bp.route("/colors/<int:color_id>/versions/new", methods=["GET", "POST"])
@login_required
def new_version(color_id):
    c = color_model.get(color_id)
    if not c:
        flash("Không tìm thấy màu này.", "error")
        return redirect(url_for("colors.index"))
    if not _can_create_version(c):
        abort(403)
    if color_version.has_pending(color_id):
        flash("Màu này đang có 1 phiên bản CHỜ DUYỆT — xử lý xong bản đó (duyệt/từ chối) trước khi tạo bản mới.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))

    pigments = pigment_model.all_ordered()
    active = color_version.get_active(color_id)
    prefill = {}
    if active:
        prefill = {r["pigment_id"]: r["ty_le"] * 100 for r in color_version.get_ratio_rows(active["id"])}

    if request.method == "POST":
        try:
            ratios = _parse_ratios(request.form, pigments)
            old_ratios = {}
            if active:
                old_ratios = {r["name"]: r["ty_le"] for r in color_version.get_ratio_rows(active["id"])}
            new_ratios_named = {p["name"]: ratios[p["id"]] for p in pigments}

            # Chỉ admin được sửa mã màu ở đây — role khác gửi gì lên cũng
            # bị bỏ qua, giữ nguyên mã màu hiện tại của màu này.
            ma_mau = request.form.get("ma_mau", c["ma_mau"]).strip() if _is_admin() else c["ma_mau"]
            color_model.update_info(
                color_id,
                ma_mau,
                request.form.get("ten_mau", c["ten_mau"] or ""),
                request.form.get("ghi_chu", c["ghi_chu"] or ""),
                hex_color=_clean_hex(request.form.get("hex_color")),
                lab=_parse_lab(request.form),
            )

            auto_approve = _can_self_approve()
            version_id = color_version.create_version(
                color_id, ratios, session["user_id"],
                ghi_chu=request.form.get("version_ghi_chu", ""),
                auto_approve=auto_approve,
            )
            v = color_version.get(version_id)
            log(session["user_id"], "TẠO PHIÊN BẢN", "color", color_id,
                {"ma_mau": ma_mau, "phien_ban": v["version_number"],
                 "cu": old_ratios, "moi": new_ratios_named,
                 "trang_thai_duyet": "DA_DUYET" if auto_approve else "CHO_DUYET"})
            if auto_approve:
                flash(f"Đã tạo phiên bản {v['version_number']} (tự động duyệt, đang dùng).", "success")
            else:
                flash(f"Đã tạo phiên bản {v['version_number']} — đang CHỜ DUYỆT.", "success")
            return redirect(url_for("colors.detail", color_id=color_id))
        except ValueError as e:
            flash(str(e), "error")
        except Exception as e:
            flash(f"Lỗi: {e}", "error")

    return render_template("colors/version_form.html", color=c, pigments=pigments, values=prefill)


@bp.route("/colors/<int:color_id>/versions/<int:version_id>/edit", methods=["GET", "POST"])
@login_required
def edit_version(color_id, version_id):
    """Sửa TRỰC TIẾP tỷ lệ của 1 phiên bản đã tồn tại — không tạo bản mới,
    không đổi trạng thái duyệt/active (trừ trường hợp resubmit — xem dưới).
    Dùng cho sửa lỗi nhỏ/tinh chỉnh bản hiện tại (khác với 'Tạo phiên bản
    mới' — dùng khi thực sự đổi công thức)."""
    c = color_model.get(color_id)
    v = color_version.get(version_id)
    if not c or not v or v["color_id"] != color_id:
        flash("Không tìm thấy phiên bản này.", "error")
        return redirect(url_for("colors.index"))
    if not _can_edit_version(v):
        abort(403)

    # Chủ sở hữu (không phải Admin/Quản lý) sửa lại bản đã bị TỪ CHỐI ->
    # coi là gửi lại xin duyệt, tự chuyển về CHỜ DUYỆT. Admin/Quản lý sửa
    # thì áp dụng ngay, giữ nguyên trạng thái hiện tại (không đổi gì).
    is_resubmit = ("duyet_mau" not in g.get("current_permissions", set())
                   and v["trang_thai_duyet"] == color_version.DUYET_TU_CHOI)

    pigments = pigment_model.all_ordered()
    prefill = {r["pigment_id"]: r["ty_le"] * 100 for r in color_version.get_ratio_rows(version_id)}

    if request.method == "POST":
        try:
            ratios = _parse_ratios(request.form, pigments)
            old_ratios_named = {p["name"]: (prefill.get(p["id"], 0) or 0) / 100.0 for p in pigments}
            new_ratios_named = {p["name"]: ratios[p["id"]] for p in pigments}
            # Chỉ admin được sửa mã màu ở đây — role khác gửi gì lên cũng
            # bị bỏ qua, giữ nguyên mã màu hiện tại của màu này.
            ma_mau = request.form.get("ma_mau", c["ma_mau"]).strip() if _is_admin() else c["ma_mau"]
            color_model.update_info(
                color_id,
                ma_mau,
                request.form.get("ten_mau", c["ten_mau"] or ""),
                request.form.get("ghi_chu", c["ghi_chu"] or ""),
                hex_color=_clean_hex(request.form.get("hex_color")),
                lab=_parse_lab(request.form),
            )
            color_version.update_ratios(
                version_id, ratios,
                ghi_chu=request.form.get("version_ghi_chu", v["ghi_chu"] or ""),
                resubmit=is_resubmit,
            )
            log(session["user_id"], "SỬA PHIÊN BẢN", "color", color_id,
                {"ma_mau": ma_mau,
                 "phien_ban": v["version_number"], "gui_lai": is_resubmit,
                 "cu": old_ratios_named, "moi": new_ratios_named})
            if is_resubmit:
                flash(f"Đã cập nhật và GỬI LẠI phiên bản {v['version_number']} — đang CHỜ DUYỆT.", "success")
            else:
                flash(f"Đã cập nhật phiên bản {v['version_number']}.", "success")
            return redirect(url_for("colors.detail", color_id=color_id))
        except ValueError as e:
            flash(str(e), "error")
        except Exception as e:
            flash(f"Lỗi: {e}", "error")

    return render_template("colors/version_edit_form.html", color=c, version=v,
                           pigments=pigments, values=prefill, is_resubmit=is_resubmit)


@bp.route("/colors/<int:color_id>/versions/<int:version_id>")
@login_required
def version_detail(color_id, version_id):
    c = color_model.get(color_id)
    v = color_version.get(version_id)
    if not c or not v or v["color_id"] != color_id:
        flash("Không tìm thấy phiên bản này.", "error")
        return redirect(url_for("colors.index"))
    d = color_version.detail(version_id)
    return render_template("colors/version_detail.html", color=c, d=d,
                           can_activate=_can_activate(),
                           can_delete_version=_can_delete_version(v),
                           can_edit_version=_can_edit_version(v))


@bp.route("/colors/<int:color_id>/versions/<int:version_id>/approve", methods=["POST"])
@permission_required("duyet_mau")
def approve_version(color_id, version_id):
    v = color_version.get(version_id)
    if not v or v["color_id"] != color_id:
        flash("Không tìm thấy phiên bản này.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    d = color_version.detail(version_id)
    if d["status"] != "ĐẠT":
        flash(f"Không thể duyệt — tổng pigment đang {d['status']}, cần đúng 100%.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    color_version.approve(version_id, session["user_id"])
    log(session["user_id"], "DUYỆT PHIÊN BẢN", "color", color_id, {"phien_ban": v["version_number"]})
    flash("Đã duyệt phiên bản — trở thành bản đang dùng.", "success")
    return redirect(url_for("colors.detail", color_id=color_id))


@bp.route("/colors/<int:color_id>/versions/<int:version_id>/reject", methods=["POST"])
@permission_required("duyet_mau")
def reject_version(color_id, version_id):
    v = color_version.get(version_id)
    if not v or v["color_id"] != color_id:
        flash("Không tìm thấy phiên bản này.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    reason = (request.form.get("ly_do") or "").strip()
    color_version.reject(version_id, session["user_id"], reason)
    log(session["user_id"], "TỪ CHỐI PHIÊN BẢN", "color", color_id,
        {"phien_ban": v["version_number"], "ly_do": reason})
    flash("Đã từ chối phiên bản.", "success")
    return redirect(url_for("colors.detail", color_id=color_id))


@bp.route("/colors/<int:color_id>/versions/<int:version_id>/activate", methods=["POST"])
@permission_required("duyet_mau")
def activate_version(color_id, version_id):
    v = color_version.get(version_id)
    if not v or v["color_id"] != color_id:
        flash("Không tìm thấy phiên bản này.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    try:
        color_version.activate(version_id)
        log(session["user_id"], "DÙNG LẠI PHIÊN BẢN", "color", color_id,
            {"phien_ban": v["version_number"]})
        flash(f"Đã đặt phiên bản {v['version_number']} làm bản đang dùng.", "success")
    except ValueError as e:
        flash(str(e), "error")
    return redirect(url_for("colors.detail", color_id=color_id))


@bp.route("/colors/<int:color_id>/versions/<int:version_id>/delete", methods=["POST"])
@login_required
def delete_version(color_id, version_id):
    v = color_version.get(version_id)
    if not v or v["color_id"] != color_id:
        flash("Không tìm thấy phiên bản này.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    if not _can_delete_version(v):
        abort(403)
    # Nếu đây là PHIÊN BẢN DUY NHẤT của màu, xóa nó = xóa luôn cả màu (xem
    # đoạn cascade bên dưới) -> phải là admin, đồng bộ với quy định "chỉ
    # admin được xóa màu". Không chặn thì Quản lý (có quyền xóa phiên bản)
    # lách qua đường này để xóa được màu.
    if len(color_version.list_for_color(color_id)) == 1 and not _is_admin():
        flash("Đây là phiên bản duy nhất của màu này — xóa sẽ xóa luôn cả màu, chỉ Admin được thực hiện.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    was_active = bool(v["is_active"])
    vnum = v["version_number"]
    color_version.delete(version_id)
    log(session["user_id"], "XÓA PHIÊN BẢN", "color", color_id, {"phien_ban": vnum})

    # Nếu vừa xóa PHIÊN BẢN CUỐI CÙNG của màu -> xóa luôn cả màu. Không làm
    # vậy sẽ để lại 1 'xác' màu 0 phiên bản: không hiện ở đâu trong giao
    # diện (thư viện cần có bản active, "Của tôi" liệt kê theo phiên bản),
    # nhưng mã màu vẫn bị chiếm giữ do ràng buộc UNIQUE — tạo lại đúng mã
    # đó sau này sẽ báo lỗi trùng mà không ai hiểu vì sao.
    if not color_version.list_for_color(color_id):
        c = color_model.get(color_id)
        ma_mau = c["ma_mau"]
        _delete_color_with_files(c)
        log(session["user_id"], "XÓA MÀU", "color", color_id,
            {"ma_mau": ma_mau, "ly_do": "Xóa phiên bản cuối cùng của màu"})
        flash(f"Đã xóa phiên bản {vnum} — đây là phiên bản cuối cùng nên màu '{ma_mau}' cũng đã được xóa luôn.", "success")
        return redirect(url_for("colors.index"))

    if was_active:
        flash(f"Đã xóa phiên bản {vnum} — hiện chưa có bản nào đang dùng, chọn 1 bản ĐÃ DUYỆT khác để dùng lại.", "success")
    else:
        flash(f"Đã xóa phiên bản {vnum}.", "success")
    return redirect(url_for("colors.detail", color_id=color_id))


@bp.route("/colors/<int:color_id>/upload-image", methods=["POST"])
@login_required
def upload_image(color_id):
    c = color_model.get(color_id)
    if not c:
        flash("Không tìm thấy màu này.", "error")
        return redirect(url_for("colors.index"))
    if not _can_manage_images(c):
        abort(403)
    files = [f for f in request.files.getlist("image") if f and f.filename]
    if not files:
        flash("Chưa chọn file.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    config.IMAGE_FOLDER.mkdir(parents=True, exist_ok=True)
    loai_anh = request.form.get("loai_anh", "Khác")
    ghi_chu = request.form.get("ghi_chu", "").strip()
    saved = 0
    for i, f in enumerate(files):
        ext = f.filename.rsplit(".", 1)[-1].lower() if "." in f.filename else ""
        if not ext or ext in config.BLOCKED_UPLOAD_EXT:
            flash(f"Bỏ qua file '{f.filename}' — đuôi .{ext or '?'} không được phép.", "error")
            continue
        filename = f"color{color_id}_{int(time.time()*1000)}_{i}_{secure_filename(Path(f.filename).stem) or 'anh'}.{ext}"
        f.save(config.IMAGE_FOLDER / filename)
        rel = f"uploads/images/{filename}"
        image_model.create(color_id, rel, loai_anh, session["user_id"], ghi_chu)
        log(session["user_id"], "TẢI ẢNH", "color", color_id, {"file": rel, "loai_anh": loai_anh})
        saved += 1
    if saved:
        flash(f"Đã tải lên {saved} file." if saved > 1 else "Đã tải file lên.", "success")
    return redirect(url_for("colors.detail", color_id=color_id))



def _delete_color_with_files(c):
    """Xóa hẳn 1 màu + dọn file ảnh/QR liên quan trên đĩa. Dùng chung cho cả
    route xóa màu trực tiếp và trường hợp xóa phiên bản cuối cùng khiến màu
    không còn phiên bản nào (tránh để lại 'xác' màu 0 phiên bản vẫn chiếm
    giữ mã màu — không hiện ở đâu trong giao diện nhưng vẫn chặn tạo lại
    đúng mã đó do ràng buộc UNIQUE)."""
    for img in image_model.list_for_color(c["id"]):
        try:
            (config.IMAGE_FOLDER / Path(img["file_path"]).name).unlink(missing_ok=True)
        except Exception:
            pass
    if c["qr_code_path"]:
        try:
            (config.QR_FOLDER / Path(c["qr_code_path"]).name).unlink(missing_ok=True)
        except Exception:
            pass
    color_model.delete(c["id"])


@bp.route("/colors/<int:color_id>/delete", methods=["POST"])
@login_required
def delete(color_id):
    c = color_model.get(color_id)
    if not c:
        flash("Không tìm thấy màu này.", "error")
        return redirect(url_for("colors.index"))
    if not _can_delete_color(c):
        abort(403)

    ma_mau = c["ma_mau"]
    _delete_color_with_files(c)
    log(session["user_id"], "XÓA MÀU", "color", color_id, {"ma_mau": ma_mau})
    flash(f"Đã xóa màu '{ma_mau}'.", "success")
    return redirect(url_for("colors.index"))


@bp.route("/colors/<int:color_id>/hex", methods=["POST"])
@permission_required("duyet_mau")
def edit_hex(color_id):
    """Sửa NHANH mã HEX + Lab đo máy ngay tại trang chi tiết — Admin/Quản
    lý (duyet_mau) mới được. Chỉ đổi màu hiển thị/Lab để so màu, KHÔNG ảnh
    hưởng công thức."""
    c = color_model.get(color_id)
    if not c:
        flash("Không tìm thấy màu này.", "error")
        return redirect(url_for("colors.index"))
    hex_moi = _clean_hex(request.form.get("hex_color"))
    if not hex_moi:
        flash("Mã HEX không hợp lệ — phải dạng #RRGGBB.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    try:
        lab_moi = _parse_lab(request.form)
    except ValueError as e:
        flash(str(e), "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    color_model.update_info(color_id, c["ma_mau"], c["ten_mau"] or "",
                            c["ghi_chu"] or "", hex_color=hex_moi, lab=lab_moi)
    log(session["user_id"], "SỬA MÃ HEX", "color", color_id,
        {"cu": c["hex_color"], "moi": hex_moi,
         "lab_cu": [c["lab_l"], c["lab_a"], c["lab_b"]] if c["lab_l"] is not None else None,
         "lab_moi": list(lab_moi) if lab_moi else None})
    flash(f"Đã đổi mã HEX → '{hex_moi}'" + (" và cập nhật Lab đo máy." if lab_moi else "."), "success")
    return redirect(url_for("colors.detail", color_id=color_id))


@bp.route("/colors/<int:color_id>/images/<int:image_id>/delete", methods=["POST"])
@login_required
def delete_image(color_id, image_id):
    img = image_model.get(image_id)
    if not img or img["color_id"] != color_id:
        flash("Không tìm thấy ảnh này.", "error")
        return redirect(url_for("colors.detail", color_id=color_id))
    # CHỈ admin được xóa ảnh — không còn cho người tự upload hay Quản lý
    # (duyet_mau) xóa nữa.
    if not _is_admin():
        abort(403)
    image_model.delete(image_id)
    try:
        (config.IMAGE_FOLDER / Path(img["file_path"]).name).unlink(missing_ok=True)
    except Exception:
        pass
    log(session["user_id"], "XÓA ẢNH", "color", color_id, {"file": img["file_path"]})
    flash("Đã xóa ảnh.", "success")
    return redirect(url_for("colors.detail", color_id=color_id))